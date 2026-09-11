#!/usr/bin/env python3
"""Local-only HU extraction. No network or mutation adapter. See input-contract.md."""
import argparse
import csv
import json
from pathlib import Path
import re
import sys
from urllib.parse import urlparse

POINTS = [1, 2, 3, 5, 8, 13]
LAYERS = {
    'CRUD': ('crud_repository', 'Persistencia, modelo y contrato CRUD; pruebas de datos.'),
    'MID': ('mid_repository', 'Reglas HU, orquestación, transformación e integración institucional.'),
    'FRONT': ('front_repository', 'Escenarios HU, interacción y consumo MID; detectar Single-SPA.'),
}
REPO = re.compile(r'[A-Za-z0-9][A-Za-z0-9-]*/[A-Za-z0-9_.-]+')


def require(condition, message):
    if not condition:
        raise ValueError(message)


def positive_int(value):
    return type(value) is int and value > 0


def clean(value):
    return '' if value is None else str(value).strip()


def valid_repo(value):
    return isinstance(value, str) and REPO.fullmatch(value) is not None and value.split('/')[1] not in {'.', '..'}


def validate_estimate(value):
    require(type(value) is int and value > 0, 'Estimate debe ser entero PH positivo, no bool/float.')
    if value > 13:
        return 'DECOMPOSE'
    require(value in POINTS, 'Estimate permitido: 1, 2, 3, 5, 8, 13.')
    return 'VALID'


def validate_config(config, run):
    require(isinstance(config, dict) and isinstance(run, dict), 'Config y run deben ser objetos JSON.')
    require(run.get('mode', 'dry-run') == 'dry-run', 'apply está deshabilitado; solo dry-run.')
    for section in ('github', 'estimation', 'project_fields', 'mapping'):
        require(isinstance(config.get(section), dict), f'{section} debe ser objeto JSON.')
    require(config.get('estimation', {}).get('allowed_points') == POINTS, 'Escala PH institucional inválida.')
    require(all(type(n) is int for n in config['estimation']['allowed_points']), 'PH no acepta booleanos.')
    require(config.get('project_fields', {}).get('developed_is_read_only') is True, 'Developed debe ser solo lectura.')
    github = config.get('github', {})
    require(github.get('organization') == 'udistrital' and type(github.get('project_number')) is int and github['project_number'] == 120, 'Esta configuración AGORA requiere udistrital / Project 120.')
    require(valid_repo(github.get('templates_repository')), 'Repositorio de plantillas inválido.')
    mapping = config.get('mapping', {})
    require(positive_int(mapping.get('header_row')), 'header_row debe ser entero positivo.')
    columns = mapping.get('columns', {})
    require(isinstance(columns, dict), 'mapping.columns debe ser objeto.')
    needed = {'requirement_id', 'description'} | {v[0] for v in LAYERS.values()}
    require(needed <= columns.keys(), 'Mapear ID, description y las tres columnas CRUD/MID/FRONT.')
    require(all(isinstance(v, str) and v.strip() == v and v for v in columns.values()), 'Encabezados mapeados deben ser texto no vacío, sin espacios externos.')
    require(len(set(columns.values())) == len(columns), 'Mapeo ambiguo: encabezado reutilizado.')
    for key in ('workbook', 'sheet', 'sprint'):
        require(isinstance(run.get(key), str) and run[key].strip(), f'Falta {key}.')
    require('assignee' not in run or isinstance(run['assignee'], str) and bool(run['assignee'].strip()), 'assignee debe ser login no vacío.')
    require('workbook_source' not in run or isinstance(run['workbook_source'], str) and bool(run['workbook_source'].strip()), 'workbook_source debe ser referencia no vacía.')
    if 'synthetic' in run:
        require(type(run['synthetic']) is bool, 'synthetic debe ser booleano.')
    return mapping


def validate_selection(selection):
    require(isinstance(selection, dict), 'selection debe ser objeto explícito.')
    keys = set(selection)
    require(keys in ({'requirement_ids'}, {'row_range'}, {'start_row', 'count'}), 'Selección ambigua: exactamente requirement_ids, row_range o start_row+count.')
    if 'requirement_ids' in selection:
        ids = selection['requirement_ids']
        require(isinstance(ids, list) and ids and all(isinstance(x, str) and x.strip() == x and x for x in ids), 'requirement_ids debe ser lista no vacía de IDs textuales.')
        require(len(set(ids)) == len(ids), 'IDs de selección duplicados.')
        return 'ids', ids
    if 'row_range' in selection:
        span = selection['row_range']
        require(isinstance(span, dict) and set(span) == {'start', 'end'}, 'row_range requiere start/end solamente.')
        start, end = span['start'], span['end']
    else:
        start, count = selection['start_row'], selection['count']
        require(positive_int(start) and positive_int(count), 'start_row/count deben ser enteros positivos.')
        end = start + count - 1
    require(positive_int(start) and positive_int(end) and end >= start, 'Rango de filas inválido.')
    return 'rows', (start, end)


def read_table(path, sheet):
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as source:
            return list(csv.reader(source)), set()
    require(path.suffix.lower() == '.xlsx', 'Solo CSV o XLSX; no macros.')
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError('XLSX requiere openpyxl 3.x; CSV funciona sin dependencias. No se instaló nada.') from exc
    # read_only avoids modifying the input. Merged cells become blanks and fail
    # required ID/header validation; selected merged non-ID data is checked via XML.
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    from posixpath import normpath
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(path) as archive:
        book_xml = ET.fromstring(archive.read('xl/workbook.xml'))
        matches = [s for s in book_xml.findall('m:sheets/m:sheet', ns) if s.attrib['name'] == sheet]
        require(len(matches) == 1, f'Hoja inexistente: {sheet}')
        rid = matches[0].attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
        rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        target = next(x.attrib['Target'] for x in rels if x.attrib['Id'] == rid)
        member = target.lstrip('/') if target.startswith('/') else normpath('xl/' + target)
        xml = ET.fromstring(archive.read(member))
        from openpyxl.utils.cell import range_boundaries
        merged_rows = set()
        for node in xml.findall('m:mergeCells/m:mergeCell', ns):
            _, lo, _, hi = range_boundaries(node.attrib['ref'])
            merged_rows.update(range(lo, hi + 1))
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        require(sheet in book.sheetnames, f'Hoja inexistente: {sheet}')
        return [list(row) for row in book[sheet].iter_rows(values_only=True)], merged_rows
    finally:
        book.close()


def prepare(config, run, base_dir=Path('.')):
    mapping = validate_config(config, run)
    strategy, requested = validate_selection(run.get('selection'))
    path = (Path(base_dir) / run['workbook']).resolve()
    table, merged_rows = read_table(path, run['sheet'])
    header_row = mapping['header_row']
    require(header_row <= len(table), 'header_row fuera del archivo.')
    require(header_row not in merged_rows, 'Encabezados combinados no admitidos.')
    headers = [clean(v) for v in table[header_row - 1]]
    require(not any(h.startswith('=') for h in headers), 'Encabezado con fórmula no admitido.')
    named = [h for h in headers if h]
    require(len(named) == len(set(named)), 'Encabezados duplicados.')
    columns = mapping['columns']
    require(set(columns.values()) <= set(headers), 'Faltan encabezados configurados.')
    id_col = headers.index(columns['requirement_id'])
    by_id, rows = {}, {}
    for number, values in enumerate(table[header_row:], header_row + 1):
        raw_identity = values[id_col] if id_col < len(values) else None
        require(raw_identity is None or isinstance(raw_identity, str), f'ID debe ser texto en fila {number}; no reconstruir formato numérico.')
        identity = clean(raw_identity)
        require(not identity.startswith('='), f'ID con fórmula en fila {number}.')
        if identity:
            require(identity not in by_id, f'ID duplicado: {identity}')
            by_id[identity] = number
        rows[number] = values
    if strategy == 'ids':
        missing = [i for i in requested if i not in by_id]
        require(not missing, f'IDs no encontrados: {missing}')
        selected = [by_id[i] for i in requested]
    else:
        start, end = requested
        require(start > header_row and end <= len(table), 'Selección fuera de filas de datos.')
        selected = list(range(start, end + 1))
    result = []
    for number in selected:
        values = rows[number]
        require(number not in merged_rows, f'Celdas combinadas en fila seleccionada {number}.')
        require(not any(clean(v).startswith('=') for v in values), f'Fórmula en fila seleccionada {number}; aportar valores confirmados.')
        require(all(not clean(v) or i < len(headers) and headers[i] for i, v in enumerate(values)), f'Datos sin encabezado en fila {number}.')
        context = {h: clean(values[i]) if i < len(values) else '' for i, h in enumerate(headers) if h}
        data = {key: context[value] for key, value in columns.items()}
        identity = data['requirement_id']
        require(identity, f'Fila {number} sin HU ID.')
        gaps = []
        if not data['description']:
            gaps.append('DESCRIPTION_MISSING: NOT_READY')
        repositories = {layer: data[column] for layer, (column, _) in LAYERS.items() if data[column]}
        require(all(valid_repo(repo) for repo in repositories.values()), f'Repositorio no canónico en {identity}.')
        if not repositories:
            gaps.append('REPOSITORIES_MISSING: NOT_READY')
        direct = data.get('hu_source_link', '')
        if direct:
            parsed = urlparse(direct)
            require(parsed.scheme in ('https', 'http') and bool(parsed.netloc), f'URL HU inválida en {identity}.')
        source = run.get('workbook_source', str(path))
        trace = {'type': 'DIRECT_LINK' if direct else 'COMPOSITE_REFERENCE', 'link': direct or None, 'workbook': source, 'sheet': run['sheet'], 'requirement_id': identity, 'row': number}
        if not direct and urlparse(source).scheme not in ('http', 'https'):
            gaps.append('ACCESSIBILITY_UNVERIFIED: validar acceso al localizador HU antes de creación futura')
        priority = data.get('priority', '')
        require(not priority or priority in ('1', '2', '3', '4'), f'Priority fuera de 1–4 en {identity}.')
        packets = []
        for layer, repo in repositories.items():
            packets.append({'local_key': f'{identity}:{layer}', 'layer': layer, 'repository': repo,
                'repository_verification': 'UNVERIFIED', 'responsibility_guide': LAYERS[layer][1],
                'hu_source': trace, 'solution_repositories': repositories,
                'siblings': [f'{identity}:{other}' for other in repositories if other != layer],
                'proposed_fields': {'Estimate': None, 'Priority': int(priority) if priority else None,
                    'Sprint': run['sprint'], 'Epic': data.get('epic') or None,
                    'Actividad': data.get('activity') or None, 'Assignee': run.get('assignee')},
                'required_analysis': ['CURRENT_TEMPLATE', 'LAYER_DESCRIPTION_AND_CRITERIA', 'DEPENDENCIES', 'PH_ESTIMATE', 'EVIDENCE_BASED_DOR', 'PROJECT_FIELD_DISCOVERY']})
        result.append({'requirement_id': identity, 'description': data['description'], 'full_context': context,
                       'hu_source': trace, 'repositories': repositories, 'packets': packets, 'gaps': gaps})
    return {'mode': 'dry-run', 'artifact': 'EXTRACTION_PACKETS_NOT_FINAL_ISSUES', 'synthetic': run.get('synthetic', False),
            'source': {'workbook': str(path), 'sheet': run['sheet']}, 'selection': run['selection'], 'requirements': result}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', type=Path, required=True)
    parser.add_argument('--run', type=Path, required=True)
    args = parser.parse_args()
    try:
        config = json.loads(args.config.read_text(encoding='utf-8'))
        run = json.loads(args.run.read_text(encoding='utf-8'))
        result = prepare(config, run, args.run.resolve().parent)
    except (ValueError, OSError, KeyError, TypeError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    sys.exit(main())
